"""Lecture des formats. Contenu FICTIF."""
import pytest

from pseudo.fichiers import FormatNonSupporte, PdfSansTexte, lire

TEXTE = "Note pour Claire Durand, jointe au 06 39 98 56 78."


def test_txt_et_md(tmp_path):
    for suffixe in (".txt", ".md"):
        f = tmp_path / f"n{suffixe}"
        f.write_text(TEXTE, encoding="utf-8")
        assert lire(f) == TEXTE


def test_docx_paragraphes_et_tableaux(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph(TEXTE)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "IBAN"
    table.rows[0].cells[1].text = "ZZ4300000000000000000000123"
    f = tmp_path / "n.docx"
    d.save(str(f))

    lu = lire(f)
    assert "Claire Durand" in lu
    # une donnee cachee dans un tableau doit etre lue, sinon elle fuiterait
    assert "ZZ4300000000000000000000123" in lu


def test_pdf_avec_texte(tmp_path):
    """Un PDF texte, meme court, doit etre lu et non pris pour un scan."""
    import shutil
    import subprocess

    if not shutil.which("cupsfilter"):
        pytest.skip("cupsfilter absent (outil macOS)")

    source = tmp_path / "s.txt"
    source.write_text(TEXTE + "\n", encoding="utf-8")
    cible = tmp_path / "s.pdf"
    with open(cible, "wb") as sortie:
        subprocess.run(["cupsfilter", str(source)], stdout=sortie,
                       stderr=subprocess.DEVNULL, check=True)

    lu = lire(cible)
    assert "Claire Durand" in lu
    assert "06 39 98 56 78" in lu


def test_pdf_scanne_refuse(tmp_path):
    """Un PDF sans texte selectionnable doit etre refuse, pas traite a vide."""
    from pypdf import PdfWriter

    ecrivain = PdfWriter()
    ecrivain.add_blank_page(width=595, height=842)
    f = tmp_path / "scan.pdf"
    with open(f, "wb") as sortie:
        ecrivain.write(sortie)

    with pytest.raises(PdfSansTexte, match="scanne"):
        lire(f)


def test_format_inconnu_refuse(tmp_path):
    f = tmp_path / "n.odt"
    f.write_text(TEXTE, encoding="utf-8")
    with pytest.raises(FormatNonSupporte):
        lire(f)
