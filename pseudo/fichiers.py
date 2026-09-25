"""Lecture des formats d'entree. Tout se fait en memoire, rien n'est recopie."""

from pathlib import Path

FORMATS = {".txt", ".md", ".docx", ".pdf"}

#: En dessous de ce nombre de caracteres par page, on considere que le PDF est
#: une image et non du texte. Le seuil est volontairement bas : un scan rend
#: zero caractere, ou une poignee d'artefacts. Un seuil plus haut refuserait a
#: tort les courriers courts, qui sont frequents.
SEUIL_PDF_SCANNE = 10


class FormatNonSupporte(ValueError):
    pass


class PdfSansTexte(ValueError):
    pass


def lire(chemin: str | Path) -> str:
    chemin = Path(chemin)
    suffixe = chemin.suffix.lower()
    if suffixe not in FORMATS:
        raise FormatNonSupporte(
            f"Format {suffixe or 'inconnu'} non pris en charge. "
            f"Formats acceptes : {', '.join(sorted(FORMATS))}."
        )
    if suffixe in (".txt", ".md"):
        return chemin.read_text(encoding="utf-8", errors="replace")
    if suffixe == ".docx":
        return _lire_docx(chemin)
    return _lire_pdf(chemin)


def _lire_docx(chemin: Path) -> str:
    import docx

    document = docx.Document(str(chemin))
    morceaux = [p.text for p in document.paragraphs]
    # Les donnees personnelles se cachent souvent dans les tableaux.
    for table in document.tables:
        for ligne in table.rows:
            morceaux.append("\t".join(cellule.text for cellule in ligne.cells))
    return "\n".join(morceaux)


def _lire_pdf(chemin: Path) -> str:
    from pypdf import PdfReader

    lecteur = PdfReader(str(chemin))
    pages = [page.extract_text() or "" for page in lecteur.pages]
    texte = "\n\n".join(pages)

    if pages and len(texte.strip()) < SEUIL_PDF_SCANNE * len(pages):
        raise PdfSansTexte(
            f"Ce PDF ne contient pas de texte selectionnable : c'est probablement "
            f"un document scanne ({len(texte.strip())} caractere(s) sur {len(pages)} page(s)). "
            f"pseudo-local ne fait pas de reconnaissance de caracteres. "
            f"Ouvrez le PDF, selectionnez le texte et collez-le directement, "
            f"ou passez le document par un outil d'OCR au prealable."
        )
    return texte
